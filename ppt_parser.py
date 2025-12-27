from pptx import Presentation
import fitz

def extract_ppt_text(path):
    if path.endswith(".pptx"):
        prs = Presentation(path)
        return "\n".join(
            shape.text for slide in prs.slides
            for shape in slide.shapes if hasattr(shape, "text")
        )
    if path.endswith(".pdf"):
        doc = fitz.open(path)
        return "\n".join(page.get_text() for page in doc)
    return ""